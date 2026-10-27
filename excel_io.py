# -*- coding: utf-8 -*-
"""Excel export/import for Wolves Sports Academy Manager (openpyxl, Arabic-safe)."""
import csv
import io
import os
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb, title, headers, rows, rtl=True):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = rtl
    ws.append(headers)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(list(r))
    # auto column widths
    for i, h in enumerate(headers, 1):
        width = len(str(h)) + 4
        for r in rows[:200]:
            v = r[i - 1] if i - 1 < len(r) else ""
            width = max(width, min(len(str(v)) + 4, 45))
        ws.column_dimensions[get_column_letter(i)].width = width
    return ws


def export_all(con):
    """Return xlsx bytes with Players / Subscriptions / Payments / Attendance / Summary sheets."""
    wb = Workbook()
    wb.remove(wb.active)

    players = con.execute(
        "SELECT p.id, p.first_name, p.middle_name, p.last_name, p.birth_date, p.school, p.gender, p.phone, "
        "p.guardian_name, p.guardian_phone, p.parent2_name, p.parent2_phone, "
        "g.name_ar AS grp, p.join_date, p.status, p.notes FROM players p LEFT JOIN groups g ON g.id=p.group_id "
        "ORDER BY p.school, p.full_name"
    ).fetchall()
    _sheet(wb, "Players اللاعبين",
           ["ID", "الاسم الأول", "اسم الأب", "اسم العائلة", "تاريخ الميلاد", "المدرسة", "الجنس", "الهاتف",
            "الأب", "هاتف الأب", "الأم", "هاتف الأم", "المجموعة", "تاريخ الانضمام", "الحالة", "ملاحظات"],
           [tuple(r) for r in players])

    subs = con.execute(
        "SELECT s.id, p.full_name, s.start_date, s.sessions_total, s.sessions_used, s.price, s.expiry_date, s.status "
        "FROM subscriptions s JOIN players p ON p.id=s.player_id ORDER BY s.start_date DESC"
    ).fetchall()
    _sheet(wb, "Subscriptions الاشتراكات",
           ["ID", "اللاعب", "تاريخ البداية", "الحصص الكلية", "الحصص المستخدمة", "السعر", "تاريخ الانتهاء", "الحالة"],
           [tuple(r) for r in subs])

    pays = con.execute(
        "SELECT pm.receipt_no, p.full_name, pm.amount, pm.date, pm.method, pm.note "
        "FROM payments pm JOIN players p ON p.id=pm.player_id ORDER BY pm.date DESC"
    ).fetchall()
    _sheet(wb, "Payments الدفعات",
           ["رقم الإيصال", "اللاعب", "المبلغ (دينار)", "التاريخ", "طريقة الدفع", "ملاحظة"],
           [tuple(r) for r in pays])

    att = con.execute(
        "SELECT a.session_date, p.full_name, g.name_ar, a.status, a.unpaid, a.marked_by "
        "FROM attendance a JOIN players p ON p.id=a.player_id LEFT JOIN groups g ON g.id=a.group_id "
        "ORDER BY a.session_date DESC"
    ).fetchall()
    _sheet(wb, "Attendance الحضور",
           ["التاريخ", "اللاعب", "المجموعة", "الحالة", "غير مدفوع", "سجّله"],
           [(r[0], r[1], r[2], {"present": "حاضر", "absent": "غائب", "excused": "معذور"}.get(r[3], r[3]),
             "نعم" if r[4] else "", r[5]) for r in att])

    # expenses — where the money went
    try:
        exp = con.execute(
            "SELECT date, COALESCE(NULLIF(category,''),'other') c, amount, note FROM expenses "
            "ORDER BY date DESC").fetchall()
    except Exception:
        exp = []
    _sheet(wb, "Expenses المصاريف",
           ["التاريخ / Date", "البند / Category", "المبلغ / Amount (JD)", "الوصف / Description"],
           [tuple(r) for r in exp])

    # coaches & salaries
    try:
        coaches = con.execute(
            "SELECT name, phone, salary_type, salary_amount, active FROM coaches ORDER BY name").fetchall()
    except Exception:
        coaches = []
    _sheet(wb, "Coaches المدربين",
           ["الاسم / Name", "الهاتف / Phone", "نوع الراتب / Salary type", "الراتب / Amount", "فعال / Active"],
           [(r[0], r[1], r[2], r[3], "نعم" if r[4] else "لا") for r in coaches])

    month = date.today().strftime("%Y-%m")
    revenue = con.execute("SELECT COALESCE(SUM(amount),0) s FROM payments WHERE date LIKE ?",
                          (month + "%",)).fetchone()["s"]
    try:
        other_exp = con.execute("SELECT COALESCE(SUM(amount),0) s FROM expenses WHERE date LIKE ?",
                                (month + "%",)).fetchone()["s"]
    except Exception:
        other_exp = 0
    kpis = [
        ("اللاعبين الفعالين / Active players",
         con.execute("SELECT COUNT(*) c FROM players WHERE status='active'").fetchone()["c"]),
        ("إيراد هذا الشهر / Revenue this month (JD)", revenue),
        ("مصاريف هذا الشهر / Expenses this month (JD)", other_exp),
        ("اشتراكات فعالة / Active subscriptions",
         con.execute("SELECT COUNT(*) c FROM subscriptions WHERE status='active'").fetchone()["c"]),
        ("حصص غير مدفوعة / Unpaid sessions",
         con.execute("SELECT COUNT(*) c FROM attendance WHERE unpaid=1").fetchone()["c"]),
        ("تاريخ التصدير / Exported at", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    _sheet(wb, "Summary الملخص", ["البند / Item", "القيمة / Value"], kpis)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


PLAYER_TEMPLATE_HEADERS = ["Full name / الاسم الكامل*", "Birth date / تاريخ الميلاد (YYYY-MM-DD)",
                           "Gender / الجنس (M/F)", "School / المدرسة", "Player phone / هاتف اللاعب",
                           "Father name / اسم الأب", "Father phone / هاتف الأب",
                           "Mother name / اسم الأم", "Mother phone / هاتف الأم",
                           "Group / المجموعة", "Notes / ملاحظات"]
EXAMPLE_MARK = "احذف هذا الصف"

# Column order used when a file has no header row (older templates / plain lists)
DEFAULT_COLS = ["name", "birth", "gender", "phone", "guardian_name", "guardian_phone", "group", "notes"]


def players_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Players"
    ws.sheet_view.rightToLeft = True
    ws.append(PLAYER_TEMPLATE_HEADERS)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    ws.append(["عمر خالد", "2013-05-01", "M", "Amman Baccalaureate School", "", "خالد أحمد", "0791234567",
               "سارة محمود", "0781234567", "", "Example / مثال - " + EXAMPLE_MARK])
    for i in range(1, len(PLAYER_TEMPLATE_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 26
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _field_for(header):
    """Which player field a header cell means (English or Arabic), or None."""
    h = (header or "").strip().lower()
    phone = any(k in h for k in ("phone", "mobile", "هاتف", "موبايل", "رقم"))
    if any(k in h for k in ("mother", "الأم", "الام", "parent 2")):
        return "parent2_phone" if phone else "parent2_name"
    if any(k in h for k in ("father", "guardian", "parent", "الأب", "الاب", "ولي")):
        return "guardian_phone" if phone else "guardian_name"
    for field, keys in (("school", ("school", "مدرس")), ("birth", ("birth", "dob", "ميلاد")),
                        ("age", ("age", "العمر")), ("gender", ("gender", "sex", "الجنس")),
                        ("group", ("group", "مجموع")), ("notes", ("note", "ملاحظ"))):
        if any(k in h for k in keys):
            return field
    if phone:
        return "phone"
    if "name" in h or "اسم" in h:
        return "name"
    return None


def _header_map(row):
    """{field: column} when the row is a header row (has a name column), else None."""
    cols = {}
    for i, cell in enumerate(row):
        f = _field_for(cell)
        if f and f not in cols:
            cols[f] = i
    return cols if "name" in cols else None


def _phone(v):
    """Excel drops the leading 0 of Jordanian mobiles (791234567 -> 0791234567)."""
    v = (v or "").strip()
    if v.endswith(".0"):
        v = v[:-2]
    if len(v) == 9 and v.startswith("7") and v.isdigit():
        v = "0" + v
    return v


def _norm_cell(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip() if v is not None else ""


def _rows_from_xlsx(file_stream):
    wb = load_workbook(file_stream, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(values_only=True):
        if row and any(c is not None and str(c).strip() for c in row):
            yield [_norm_cell(c) for c in row]


def _rows_from_delimited(file_stream, delimiter):
    raw = file_stream.read()
    if isinstance(raw, bytes):
        # handle a UTF-8 BOM from Excel-exported CSVs
        raw = raw.decode("utf-8-sig", errors="replace")
    lines = [ln for ln in raw.splitlines() if ln.strip()]
    for r in csv.reader(lines, delimiter=delimiter):
        if r and any(c.strip() for c in r):
            yield [c.strip() for c in r]


def _rows_from_txt(file_stream):
    raw = file_stream.read()
    if isinstance(raw, bytes):
        raw = raw.decode("utf-8-sig", errors="replace")
    for ln in raw.splitlines():
        name = ln.strip()
        if name:
            yield [name]


def import_players(con, file_storage, suggest_group_fn, filename=None):
    """Bulk-load players from .xlsx / .csv / .tsv / .txt.

    Columns are matched by their header (English or Arabic), in any order: name,
    birth date or age, gender, school, player phone, father name/phone, mother
    name/phone, group, notes. Files without a header row use DEFAULT_COLS.
    TXT imports names only. Returns (imported, skipped, errors).
    """
    filename = (filename or getattr(file_storage, "filename", "") or "").lower()
    stream = getattr(file_storage, "stream", file_storage)

    # sniff the leading bytes so a raw BytesIO (no filename) is still detected
    magic = b""
    try:
        pos = stream.tell()
        head = stream.read(4)
        stream.seek(pos)
        magic = head if isinstance(head, bytes) else head.encode("utf-8", "replace")
    except Exception:
        pass
    is_xlsx = filename.endswith(".xlsx") or magic[:2] == b"PK"

    if is_xlsx:
        row_iter = _rows_from_xlsx(stream)
    elif filename.endswith(".tsv"):
        row_iter = _rows_from_delimited(stream, "\t")
    elif filename.endswith(".txt"):
        row_iter = _rows_from_txt(stream)
    else:
        # csv (default for anything text-like, incl. unknown/no extension)
        row_iter = _rows_from_delimited(stream, ",")

    imported, skipped, errors = 0, 0, []
    groups = {}
    for r in con.execute("SELECT id, name_ar, name_en FROM groups").fetchall():
        groups[(r["name_ar"] or "").strip()] = r["id"]
        groups[(r["name_en"] or "").strip().lower()] = r["id"]

    cols = None
    for i, row in enumerate(row_iter, start=1):
        if cols is None:
            cols = _header_map(row)
            if cols:
                continue  # header row
            cols = {f: n for n, f in enumerate(DEFAULT_COLS)}
        g = lambda f: (row[cols[f]].strip() if f in cols and len(row) > cols[f] and row[cols[f]] else "")
        name = " ".join(g("name").split())
        if not name or any(EXAMPLE_MARK in c for c in row):
            continue
        try:
            birth = g("birth")[:10]
            age = g("age").split(".")[0]
            if not birth and age.isdigit():
                birth = f"{date.today().year - int(age)}-01-01"
            gd = g("gender").upper()
            gender = "F" if gd.startswith("F") or "أنثى" in gd or "بنت" in gd else "M"
            grp_name = g("group")
            group_id = groups.get(grp_name) or groups.get(grp_name.lower()) or suggest_group_fn(con, birth, gender)
            if con.execute("SELECT id FROM players WHERE full_name=?", (name,)).fetchone():
                skipped += 1
                continue
            parts = name.split()
            first = parts[0] if parts else ""
            last = parts[-1] if len(parts) > 1 else ""
            middle = " ".join(parts[1:-1])
            con.execute(
                "INSERT INTO players (full_name,first_name,middle_name,last_name,birth_date,gender,school,phone,"
                "guardian_name,guardian_phone,parent2_name,parent2_phone,group_id,join_date,notes,status) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'active')",
                (name, first, middle, last, birth, gender, g("school"), _phone(g("phone")),
                 g("guardian_name"), _phone(g("guardian_phone")), g("parent2_name"), _phone(g("parent2_phone")),
                 group_id, date.today().isoformat(), g("notes")),
            )
            imported += 1
        except Exception as e:
            errors.append(f"row {i}: {e}")
    con.commit()
    return imported, skipped, errors


def weekly_auto_export(con, exports_dir):
    """Write a dated full export to /exports if none exists in the last 7 days."""
    os.makedirs(exports_dir, exist_ok=True)
    existing = sorted(f for f in os.listdir(exports_dir) if f.startswith("academy-export-"))
    if existing:
        last = existing[-1][len("academy-export-"):len("academy-export-") + 10]
        try:
            if (date.today() - date.fromisoformat(last)).days < 7:
                return None
        except ValueError:
            pass
    path = os.path.join(exports_dir, f"academy-export-{date.today().isoformat()}.xlsx")
    with open(path, "wb") as f:
        f.write(export_all(con))
    return path
